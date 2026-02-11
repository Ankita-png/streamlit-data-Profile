"""
Time Management Excel Generator Agent
This module provides functionality to generate time management Excel files
with various templates and customization options.
"""

import pandas as pd
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io


class TimeManagementAgent:
    """
    Agent class for generating time management Excel files
    with different templates and configurations.
    """
    
    def __init__(self, template_type='daily'):
        """
        Initialize the Time Management Agent
        
        Args:
            template_type (str): Type of template ('daily', 'weekly', 'monthly')
        """
        self.template_type = template_type
        self.workbook = None
        
    def generate_daily_template(self, start_date=None, num_days=7):
        """
        Generate a daily time management template
        
        Args:
            start_date (datetime): Starting date for the template
            num_days (int): Number of days to include
            
        Returns:
            io.BytesIO: Excel file in memory
        """
        if start_date is None:
            start_date = datetime.now()
            
        # Create data for daily template
        time_slots = []
        for hour in range(6, 23):  # 6 AM to 10 PM
            for minute in [0, 30]:
                time_slots.append(f"{hour:02d}:{minute:02d}")
        
        # Create DataFrame with time slots
        data = {'Time': time_slots}
        for i in range(num_days):
            date = start_date + timedelta(days=i)
            date_str = date.strftime('%Y-%m-%d (%A)')
            data[date_str] = [''] * len(time_slots)
        
        df = pd.DataFrame(data)
        
        # Create Excel file with styling
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Daily Schedule', index=False)
            workbook = writer.book
            worksheet = writer.sheets['Daily Schedule']
            
            # Apply styling
            self._apply_styling(worksheet, 'Daily Time Management')
            
        output.seek(0)
        return output
    
    def generate_weekly_template(self, start_date=None, num_weeks=4):
        """
        Generate a weekly time management template
        
        Args:
            start_date (datetime): Starting date for the template
            num_weeks (int): Number of weeks to include
            
        Returns:
            io.BytesIO: Excel file in memory
        """
        if start_date is None:
            start_date = datetime.now()
            # Adjust to start of week (Monday)
            start_date = start_date - timedelta(days=start_date.weekday())
        
        # Create data for weekly template
        days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
        
        weekly_data = []
        for week in range(num_weeks):
            week_start = start_date + timedelta(weeks=week)
            week_data = {
                'Week': f"Week {week + 1} ({week_start.strftime('%Y-%m-%d')})",
                'Tasks': '',
                'Priority': '',
                'Status': '',
                'Hours Allocated': '',
                'Notes': ''
            }
            weekly_data.append(week_data)
            
            # Add daily breakdown
            for day in days:
                day_data = {
                    'Week': f"  {day}",
                    'Tasks': '',
                    'Priority': '',
                    'Status': '',
                    'Hours Allocated': '',
                    'Notes': ''
                }
                weekly_data.append(day_data)
        
        df = pd.DataFrame(weekly_data)
        
        # Create Excel file with styling
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Weekly Schedule', index=False)
            workbook = writer.book
            worksheet = writer.sheets['Weekly Schedule']
            
            # Apply styling
            self._apply_styling(worksheet, 'Weekly Time Management')
            
        output.seek(0)
        return output
    
    def generate_monthly_template(self, start_date=None, num_months=3):
        """
        Generate a monthly time management template
        
        Args:
            start_date (datetime): Starting date for the template
            num_months (int): Number of months to include
            
        Returns:
            io.BytesIO: Excel file in memory
        """
        if start_date is None:
            start_date = datetime.now().replace(day=1)
        
        # Create data for monthly template
        monthly_data = []
        
        for month_offset in range(num_months):
            # Calculate month
            target_month = start_date.month + month_offset
            target_year = start_date.year + (target_month - 1) // 12
            target_month = ((target_month - 1) % 12) + 1
            
            month_date = datetime(target_year, target_month, 1)
            month_name = month_date.strftime('%B %Y')
            
            # Add month header
            monthly_data.append({
                'Date': month_name,
                'Goal/Task': '',
                'Category': '',
                'Priority': '',
                'Deadline': '',
                'Progress': '',
                'Notes': ''
            })
            
            # Add weeks in month
            for week in range(1, 5):
                monthly_data.append({
                    'Date': f"  Week {week}",
                    'Goal/Task': '',
                    'Category': '',
                    'Priority': '',
                    'Deadline': '',
                    'Progress': '',
                    'Notes': ''
                })
        
        df = pd.DataFrame(monthly_data)
        
        # Create Excel file with styling
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Monthly Plan', index=False)
            workbook = writer.book
            worksheet = writer.sheets['Monthly Plan']
            
            # Apply styling
            self._apply_styling(worksheet, 'Monthly Time Management')
            
        output.seek(0)
        return output
    
    def generate_task_tracker(self):
        """
        Generate a comprehensive task tracker template
        
        Returns:
            io.BytesIO: Excel file in memory
        """
        # Create task tracker data
        task_data = {
            'Task ID': [f'T{i:03d}' for i in range(1, 21)],
            'Task Name': [''] * 20,
            'Description': [''] * 20,
            'Priority': [''] * 20,
            'Status': [''] * 20,
            'Start Date': [''] * 20,
            'Due Date': [''] * 20,
            'Assigned To': [''] * 20,
            'Time Estimate (hrs)': [''] * 20,
            'Time Spent (hrs)': [''] * 20,
            'Completion %': [''] * 20,
            'Notes': [''] * 20
        }
        
        df = pd.DataFrame(task_data)
        
        # Create Excel file with styling
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Task Tracker', index=False)
            workbook = writer.book
            worksheet = writer.sheets['Task Tracker']
            
            # Apply styling
            self._apply_styling(worksheet, 'Task Tracker')
            
            # Add data validation for Priority and Status columns
            from openpyxl.worksheet.datavalidation import DataValidation
            
            # Priority validation
            priority_dv = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
            worksheet.add_data_validation(priority_dv)
            priority_dv.add(f'D2:D{len(df) + 1}')
            
            # Status validation
            status_dv = DataValidation(type="list", formula1='"Not Started,In Progress,Completed,On Hold"', allow_blank=True)
            worksheet.add_data_validation(status_dv)
            status_dv.add(f'E2:E{len(df) + 1}')
            
        output.seek(0)
        return output
    
    def _apply_styling(self, worksheet, title):
        """
        Apply consistent styling to the worksheet
        
        Args:
            worksheet: openpyxl worksheet object
            title (str): Title for the worksheet
        """
        # Define styles
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Style header row
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border
        
        # Auto-adjust column widths
        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
        
        # Freeze header row
        worksheet.freeze_panes = 'A2'


def generate_time_management_excel(template_type='daily', **kwargs):
    """
    Convenience function to generate time management Excel files
    
    Args:
        template_type (str): Type of template ('daily', 'weekly', 'monthly', 'task_tracker')
        **kwargs: Additional arguments for specific template types
        
    Returns:
        io.BytesIO: Excel file in memory
    """
    agent = TimeManagementAgent(template_type)
    
    if template_type == 'daily':
        return agent.generate_daily_template(
            start_date=kwargs.get('start_date'),
            num_days=kwargs.get('num_days', 7)
        )
    elif template_type == 'weekly':
        return agent.generate_weekly_template(
            start_date=kwargs.get('start_date'),
            num_weeks=kwargs.get('num_weeks', 4)
        )
    elif template_type == 'monthly':
        return agent.generate_monthly_template(
            start_date=kwargs.get('start_date'),
            num_months=kwargs.get('num_months', 3)
        )
    elif template_type == 'task_tracker':
        return agent.generate_task_tracker()
    else:
        raise ValueError(f"Unknown template type: {template_type}")
